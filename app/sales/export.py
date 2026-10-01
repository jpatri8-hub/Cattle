"""Builds the Excel export of sales within a date range."""
import openpyxl
from openpyxl.styles import Font

from app.models import Sale

COLUMN_HEADERS = [
    "Invoice", "Sale Date", "Category", "Buyer", "Animal ID", "Animal Type",
    "Price", "Weight", "Payment Method", "Notes",
]


def sales_in_range(start_date, end_date):
    return (
        Sale.query.filter(Sale.sale_date >= start_date, Sale.sale_date <= end_date)
        .order_by(Sale.sale_date, Sale.id)
        .all()
    )


def build_sales_workbook(start_date, end_date):
    sales = sales_in_range(start_date, end_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"

    bold = Font(bold=True)
    for col_idx, header in enumerate(COLUMN_HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header).font = bold

    row_idx = 2
    total_price = 0.0
    line_count = 0
    for sale in sales:
        for line in sale.lines:
            animal = line.animal
            ws.cell(row=row_idx, column=1, value=sale.invoice_number)
            ws.cell(row=row_idx, column=2, value=sale.sale_date.isoformat())
            ws.cell(row=row_idx, column=3, value=sale.category.name if sale.category else "")
            ws.cell(row=row_idx, column=4, value=sale.buyer.name if sale.buyer else "")
            ws.cell(row=row_idx, column=5, value=animal.display_id if animal else "")
            ws.cell(row=row_idx, column=6, value=animal.animal_type.name if animal and animal.animal_type else "")
            ws.cell(row=row_idx, column=7, value=float(line.price))
            ws.cell(row=row_idx, column=8, value=float(line.weight) if line.weight else None)
            ws.cell(row=row_idx, column=9, value=sale.payment_method or "")
            ws.cell(row=row_idx, column=10, value=sale.notes or "")
            total_price += float(line.price)
            line_count += 1
            row_idx += 1

    ws.cell(row=row_idx, column=4, value="Total").font = bold
    ws.cell(row=row_idx, column=5, value=line_count).font = bold
    ws.cell(row=row_idx, column=7, value=round(total_price, 2)).font = bold

    for col_idx in range(1, len(COLUMN_HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 16

    return wb
