"""American Angus Association registration report.

Builds the "Registrations" workbook the ranch submits for its Registered
Angus calves, following the ranch's column-by-column instructions for the
AAA template. Calving seasons run September-April and are reported the
following May-October, so the report is keyed by "season end year" - the
2026-2027 season (Sept 2026 - Apr 2027) is season_end_year=2027.

Brand numbers are not assigned at birth - they're only assigned to calves
the user actually selects to submit for registration, so the sequence
reflects real submissions rather than every calf ever born.
"""
from datetime import date

import openpyxl
from openpyxl.styles import Font

from app.lifecycle import generate_brand_number
from app.models import Animal, AnimalType, CalfRecord, SEX_MALE

MEMBER_CODE = "724338"
REGISTERED_CALF_TYPES = ["Registered Angus Bull Calf", "Registered Angus Heifer Calf"]

COLUMN_HEADERS = [
    "CALF TAG", "SEX*", "BIRTH DATE*", "ANGUS NAME*", "PRIMARY ID*", "TATTOO/ BRAND*",
    "840 EID", "SECONDARY ID", "ARTIFICIAL INSEMINATION?*", "TWIN INDICATOR*", "SIRE REG*",
    "SIRE NAME", "DAM TAG", "DAM TATTOO", "DAM REG*", "DAM NAME", "FIRST OWNER*",
    "BULL PERMIT", "PERMIT TYPE", "EMBRYO TRANSPLANT?*", "IVF?", "EMBRYO REMOVAL DATE",
    "STORE ELECTRONICALLY?*", "BIRTH WEIGHT", "BIRTH GROUP CODE", "CALVING EASE",
]


def season_bounds(season_end_year):
    start = date(season_end_year - 1, 9, 1)
    end = date(season_end_year, 4, 30)
    return start, end


def default_season_end_year(today=None):
    today = today or date.today()
    return today.year if today.month <= 10 else today.year + 1


def _first_sire(calf):
    if calf.sire_id and calf.sire:
        return calf.sire
    if calf.candidate_sires:
        return sorted(calf.candidate_sires, key=lambda b: b.tag_id or b.temp_id or "")[0]
    return None


def eligible_calf_records(season_end_year):
    """Every Registered Angus calf born in the given season, oldest first -
    the full pool the registration report can select from, regardless of
    whether a brand number has been assigned yet."""
    start, end = season_bounds(season_end_year)
    return (
        CalfRecord.query.join(Animal, CalfRecord.calf_animal_id == Animal.id)
        .join(AnimalType, Animal.animal_type_id == AnimalType.id)
        .filter(
            CalfRecord.calving_date >= start, CalfRecord.calving_date <= end,
            AnimalType.name.in_(REGISTERED_CALF_TYPES),
        )
        .order_by(CalfRecord.calving_date, CalfRecord.dam_id)
        .all()
    )


def _dam_counts(calf_records):
    counts = {}
    for c in calf_records:
        counts[c.dam_id] = counts.get(c.dam_id, 0) + 1
    return counts


def _row_for_calf(c, dam_counts):
    calf = c.calf_animal
    dam = c.dam
    if not calf or not dam:
        return None
    sire = _first_sire(calf)
    return {
        "CALF TAG": calf.display_id,
        "SEX*": "Bull" if calf.sex == SEX_MALE else "Heifer",
        "BIRTH DATE*": c.calving_date,
        "ANGUS NAME*": "",
        "PRIMARY ID*": calf.brand_number or "",
        "TATTOO/ BRAND*": "",
        "840 EID": "",
        "SECONDARY ID": "V",
        "ARTIFICIAL INSEMINATION?*": "",
        "TWIN INDICATOR*": 0 if dam_counts.get(dam.id) == 1 else "",
        "SIRE REG*": sire.registration_number if sire and sire.registration_number else "",
        "SIRE NAME": "",
        "DAM TAG": dam.display_id,
        "DAM TATTOO": "",
        "DAM REG*": dam.registration_number or "",
        "DAM NAME": "",
        "FIRST OWNER*": MEMBER_CODE,
        "BULL PERMIT": "",
        "PERMIT TYPE": "",
        "EMBRYO TRANSPLANT?*": "",
        "IVF?": "",
        "EMBRYO REMOVAL DATE": "",
        "STORE ELECTRONICALLY?*": "Y",
        "BIRTH WEIGHT": float(c.birth_weight) if c.birth_weight else "",
        "BIRTH GROUP CODE": "",
        "CALVING EASE": c.calving_ease or "",
    }


def registration_rows_for_selection(season_end_year, selected_ids=None):
    """Rows for the given season, optionally limited to selected_ids (a set
    of CalfRecord ids). Twin indicators are always computed from the whole
    season's calves so they stay correct even when exporting a subset."""
    calves = eligible_calf_records(season_end_year)
    dam_counts = _dam_counts(calves)
    if selected_ids is not None:
        calves = [c for c in calves if c.id in selected_ids]
    rows = []
    for c in calves:
        row = _row_for_calf(c, dam_counts)
        if row:
            rows.append(row)
    return rows


def registration_report_rows(season_end_year):
    return registration_rows_for_selection(season_end_year)


def assign_brand_numbers(calf_records):
    """Assigns a brand number to each given CalfRecord's calf that doesn't
    already have one, oldest calving date first, so a partial/re-export never
    changes a previously-assigned number."""
    assigned = []
    for c in sorted(calf_records, key=lambda r: r.calving_date):
        calf = c.calf_animal
        if calf and not calf.brand_number:
            calf.brand_number = generate_brand_number(c.calving_date)
            calf.name = calf.brand_number
            assigned.append(calf)
    return assigned


def build_registration_workbook(season_end_year, selected_ids=None):
    rows = registration_rows_for_selection(season_end_year, selected_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registrations"

    ws.cell(row=1, column=1, value=f"Member Code: {MEMBER_CODE}")
    ws.cell(row=2, column=1, value="REGISTRATIONS")
    ws.cell(row=2, column=7, value=date.today().strftime("%m/%d/%y"))
    ws.cell(row=2, column=8, value=f"Total Animals:  {len(rows)}")

    bold = Font(bold=True)
    for col_idx, header in enumerate(COLUMN_HEADERS, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = bold

    for row_idx, row in enumerate(rows, start=4):
        for col_idx, header in enumerate(COLUMN_HEADERS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row[header])

    for col_idx in range(1, len(COLUMN_HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 16

    return wb
